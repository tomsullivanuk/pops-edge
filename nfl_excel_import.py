"""Read-only ELWAY workbook ingestion and explicit per-week confirmation."""
from collections import Counter
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile
import openpyxl
import nfl_forecast_import as base
import nfl_forecast_time as timing

VERSION='elway-excel-import-v1'
MAX_BYTES=8*1024*1024


def parse(raw, file_time=None):
    if len(raw)>MAX_BYTES:raise ValueError('Workbook exceeds 8 MB')
    with ZipFile(BytesIO(raw)) as z:
        if sum(i.file_size for i in z.infolist())>40*1024*1024:raise ValueError('Expanded workbook exceeds limit')
    wb=openpyxl.load_workbook(BytesIO(raw),data_only=False,read_only=True,keep_links=False)
    try:
        if len(wb.worksheets)!=1:raise ValueError('Use one worksheet containing the complete copied table')
        ws=wb.active
        if ws.max_row>1000 or ws.max_column>30:raise ValueError('Unexpected workbook dimensions')
        cells=list(ws.values)
        if any(isinstance(v,str) and v.startswith('=') for row in cells for v in row):raise ValueError('Paste source values, not formulas')
        text='\n'.join(str(v) for row in cells for v in row if v is not None)
        season=re.findall(r'every (\d{4}) regular-season game',text)
        updated=re.findall(r'Updated (\w+ \d+, \d{4}) at (\d+:\d+ [AP]M) (EDT|EST)',text)
        counts=re.findall(r'(\d+) games',re.sub(r'Week \d{1,2}', '',text))
        # Preserve legacy rejection text as well as values for saved-report replay.
        if file_time is None and (len(season)!=1 or len(updated)!=1 or len(counts)!=1):
            raise ValueError('Include heading, game count and published update time')
        if len(season)!=1 or len(counts)!=1:raise ValueError('Include heading and complete game count')
        missing = text.splitlines().count('Updated time unavailable') == 1
        proxy = not updated and missing and len([line for line in text.splitlines() if line.startswith('Updated ')]) == 1
        if len(updated)==1 and (not missing or file_time is None):
            date,time,tz=updated[0]
            at=datetime.strptime(date+' '+time,'%B %d, %Y %I:%M %p').isoformat()+('-04:00' if tz=='EDT' else '-05:00')
        elif proxy and file_time is not None:
            at=timing.validate(file_time,raw,file_time['observed_at'])
            if base.timestamp(at).year not in (int(season[0]),int(season[0])+1):
                raise ValueError('File creation year conflicts with forecast season')
        else:raise ValueError('ELWAY published update time is missing or ambiguous; file creation evidence is required for Updated time unavailable')
        headers=[i for i,r in enumerate(cells) if list(r[:9])==['Wk','Home','Avg.','Win','Away','Avg.','Win','Home','Total']]
        if len(headers)!=1:raise ValueError('Copied table headers not recognized')
        start=headers[0]+2;rows=[];seen=set();ended=False
        for index,r in enumerate(cells[start:],start+1):
            if isinstance(r[0],str) and r[0].startswith('Note:'):ended=True;break
            if not any(v is not None for v in r):continue
            m=re.fullmatch(r'(\d{1,2})(N)?(\*)?',str(r[0]).replace(' ',''))
            if not m or len(r)<9 or any(v is None for v in r[:9]):raise ValueError(f'Incomplete or unrecognized game at Excel row {index}')
            week=int(m[1]);home,away=r[1],r[4]
            if not 1<=week<=18 or home not in base.TEAMS or away not in base.TEAMS or home==away:raise ValueError(f'Invalid week/teams at row {index}')
            for team in (home,away):
                if (week,team) in seen:raise ValueError(f'Duplicate team in Week {week}')
                seen.add((week,team))
            probs=[]
            for value in (r[3],r[6]):
                if type(value) not in (int,float):raise ValueError(f'Expected numeric Excel percentage at row {index}')
                p=Decimal(str(value))
                if not 0<=p<=1 or p*1000!=(p*1000).to_integral_value():raise ValueError('Expected ELWAY probabilities in tenths of one percent')
                probs.append(p)
            if sum(probs)>1:raise ValueError('Win probabilities exceed 100%')
            for v in (r[2],r[5],r[8]):
                if type(v) not in (int,float) or not Decimal(str(v)).is_finite() or v<0:raise ValueError('Invalid score/total')
            if r[7]!='PK' and (type(r[7]) not in (int,float) or not Decimal(str(r[7])).is_finite()):raise ValueError('Invalid spread')
            rows.append(dict(week=week,home=home,away=away,home_win=f'{probs[0]*100:.1f}%',away_win=f'{probs[1]*100:.1f}%',neutral=bool(m[2]),conditional=bool(m[3]),excel_row=index,home_points=str(r[2]),away_points=str(r[5]),spread=str(r[7]),total=str(r[8])))
        if not ended or len(rows)!=int(counts[0]) or not 1<=len(rows)<=272:raise ValueError('Copied game count or footnotes are incomplete')
        if len(rows)==272 and (len({r['week'] for r in rows})!=18 or set(Counter(t for r in rows for t in (r['home'],r['away'])).values())!={17}):raise ValueError('Full-season coverage does not reconcile')
        return dict(season=int(season[0]),updated_at=at,rows=rows,game_count=len(rows),weeks=sorted({r['week'] for r in rows}),neutral_games=sum(r['neutral'] for r in rows),**({'file_time':file_time} if proxy else {}))
    finally:wb.close()


def prepare(raw,name,store,file_time=None):
    store=Path(store);identity=base.digest(raw);folder=store/'sources'/identity
    receipt=dict(schema=VERSION,source_sha256=identity,original_name=Path(name).name,imported_at=base.now())
    if (folder/'receipt.json').exists():receipt=__import__('json').loads((folder/'receipt.json').read_text())
    base.write_once(folder/'source.xlsx',raw);base.write_once(folder/'receipt.json',base.encode(receipt))
    time_path=folder/'file-time.json'
    if time_path.exists():file_time=__import__('json').loads(time_path.read_text())
    result=parse(raw,file_time)
    if result.get('file_time'):timing.validate(file_time,raw,base.now())
    if base.timestamp(result['updated_at'])>base.timestamp(receipt['imported_at']):raise ValueError('Model update time is after import')
    if result.get('file_time'):base.write_once(time_path,base.encode(result['file_time']))
    return dict(source=receipt,**result)


def verify(candidate,week,reviewer,store,confirmed=False):
    if confirmed is not True or not isinstance(reviewer,str) or not reviewer.strip():raise ValueError('Confirm the selected week preview and reviewer name')
    receipt=candidate['source'];raw=(Path(store)/'sources'/receipt['source_sha256']/'source.xlsx').read_bytes();parsed=parse(raw)
    if week not in parsed['weeks']:raise ValueError('Week is absent from workbook')
    chosen=[r for r in parsed['rows'] if r['week']==week]
    if any(r['conditional'] for r in chosen):raise ValueError('Conditional forecast rows need a separate review; this week cannot be activated')
    review=dict(source_sha256=receipt['source_sha256'],season=parsed['season'],week=week,updated_at=parsed['updated_at'],expected_games=len(chosen),metadata_reviewed=True,conditional_markers_checked=True,rows=[dict(r,reviewed=True) for r in chosen])
    rows=base.validate(review,receipt)
    key=base.digest(base.encode(dict(review=review,reviewer=reviewer,supersedes=None)))
    record=dict(schema=VERSION,source=receipt,review=review,rows=rows,reviewer=reviewer,supersedes=None,verification_id=key,verified_at=base.now(),scope='Owner-confirmed Excel transcription for selected week; not screenshot verification')
    path=Path(store)/'verified'/f'{key}.json'
    if path.exists():return path
    validate(record,raw,receipt);base.write_once(path,base.encode(record));return path


def validate(record,raw,receipt):
    if record['schema']!=VERSION or record['source']!=receipt or base.digest(raw)!=receipt['source_sha256']:raise ValueError('Workbook source identity mismatch')
    parsed=parse(raw);r=record['review'];chosen=[x for x in parsed['rows'] if x['week']==r['week']]
    expected=dict(source_sha256=receipt['source_sha256'],season=parsed['season'],week=r['week'],updated_at=parsed['updated_at'],expected_games=len(chosen),metadata_reviewed=True,conditional_markers_checked=True,rows=[dict(x,reviewed=True) for x in chosen])
    if r!=expected or any(x['conditional'] for x in chosen) or not record['reviewer'].strip():raise ValueError('Workbook confirmation mismatch')
    key=base.digest(base.encode(dict(review=r,reviewer=record['reviewer'],supersedes=None)))
    if record['verification_id']!=key or record['rows']!=base.validate(r,receipt):raise ValueError('Workbook verification mismatch')
    if base.timestamp(record['verified_at'])<base.timestamp(receipt['imported_at']):raise ValueError('Workbook chronology mismatch')
    return record

AUTO_VERSION='elway-excel-validation-v2'


def automatic_rows(parsed,week,receipt):
    chosen=[r for r in parsed['rows'] if r['week']==week]
    if not chosen or any(r['conditional'] for r in chosen):raise ValueError(f'Week {week}: conditional or missing rows need attention')
    # Reuse numeric/identity checks; these temporary flags are not human attestations
    # and are never persisted in the machine-validation record.
    shape=dict(source_sha256=receipt['source_sha256'],season=parsed['season'],week=week,updated_at=parsed['updated_at'],expected_games=len(chosen),metadata_reviewed=True,rows=[dict(r,reviewed=True) for r in chosen])
    return base.validate(shape,receipt)


def validate_automatically(candidate,week,store):
    receipt=candidate['source'];raw=(Path(store)/'sources'/receipt['source_sha256']/'source.xlsx').read_bytes();parsed=parse(raw,candidate.get('file_time'))
    review=dict(source_sha256=receipt['source_sha256'],season=parsed['season'],week=week,updated_at=parsed['updated_at'],expected_games=sum(r['week']==week for r in parsed['rows']))
    if parsed.get('file_time'):review['file_time']=parsed['file_time']
    rows=automatic_rows(parsed,week,receipt)
    key=base.digest(base.encode(dict(schema=AUTO_VERSION,source=receipt,review=review,rows=rows)))
    record=dict(schema=AUTO_VERSION,source=receipt,review=review,rows=rows,verification_id=key,verified_at=base.now(),scope='Automated workbook validation; no human review attested')
    path=Path(store)/'verified'/f'{key}.json'
    if not path.exists():validate_auto(record,raw,receipt);base.write_once(path,base.encode(record))
    else:validate_auto(__import__('json').loads(path.read_text()),raw,receipt)
    return path


def validate_auto(record,raw,receipt):
    if record['schema']!=AUTO_VERSION or record['source']!=receipt or base.digest(raw)!=receipt['source_sha256']:raise ValueError('Workbook source identity mismatch')
    file_time=record['review'].get('file_time')
    if file_time:timing.validate(file_time,raw,record['verified_at'])
    parsed=parse(raw,file_time);week=record['review']['week']
    review=dict(source_sha256=receipt['source_sha256'],season=parsed['season'],week=week,updated_at=parsed['updated_at'],expected_games=sum(r['week']==week for r in parsed['rows']))
    if parsed.get('file_time'):review['file_time']=parsed['file_time']
    rows=automatic_rows(parsed,week,receipt)
    key=base.digest(base.encode(dict(schema=AUTO_VERSION,source=receipt,review=review,rows=rows)))
    if record['review']!=review or record['rows']!=rows or record['verification_id']!=key or record.get('reviewer') is not None:raise ValueError('Automated validation mismatch')
    if base.timestamp(record['verified_at'])<base.timestamp(receipt['imported_at']):raise ValueError('Workbook chronology mismatch')
    return record
